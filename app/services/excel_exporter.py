from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from app.models.cccd_result import CCCDResult

HEADERS = [
    "STT",
    "Tên ảnh",
    "Số CCCD",
    "Số CMND cũ",
    "Họ và tên",
    "Ngày sinh",
    "Giới tính",
    "Nơi cư trú",
    "Ngày cấp",
    "Họ tên cha",
    "Họ tên mẹ",
    "Nguồn nhận dạng",
    "Kết quả",
    "Ghi chú",
]

TEXT_COLUMNS = {3, 4}  # 1-based: Số CCCD, Số CMND cũ


def export_cccd_results(path: str | Path, results: list[CCCDResult]) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kết quả CCCD"
    sheet.append(HEADERS)

    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.font = header_font

    for index, result in enumerate(results, start=1):
        values = [index, *result.display_row()]
        sheet.append(values)
        row_idx = index + 1
        for col in TEXT_COLUMNS:
            cell = sheet.cell(row=row_idx, column=col)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [6, 28, 16, 14, 28, 14, 12, 36, 14, 24, 24, 16, 18, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        for cell in sheet[get_column_letter(index)]:
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    workbook.save(destination)
    workbook.close()
    return destination
