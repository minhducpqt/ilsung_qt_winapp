from __future__ import annotations

from pathlib import Path


def read_xlsx(path: str | Path) -> list[list[object]]:
    """Read the active sheet of an .xlsx file into rows."""
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def write_xlsx(path: str | Path, rows: list[list[object]], sheet_name: str = "Sheet1") -> Path:
    """Create an .xlsx file from rows."""
    from openpyxl import Workbook

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(list(row))
    workbook.save(destination)
    workbook.close()
    return destination
