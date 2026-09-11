from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.features.cccd_pairing.models import CCCDPersonPairRecord

HEADERS = [
    "STT",
    "Số CCCD",
    "Số CMND cũ",
    "Họ và tên",
    "Ngày sinh",
    "Giới tính",
    "Địa chỉ",
    "Ngày cấp",
    "Họ tên cha",
    "Họ tên mẹ",
    "File mặt trước",
    "File mặt sau",
    "Trạng thái nhận dạng",
    "Trạng thái ghép",
    "Ghi chú",
]
TEXT_COLUMNS = {2, 3}


def export_pairing_xlsx(path: str | Path, persons: list[CCCDPersonPairRecord]) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Danh sách CCCD"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, person in enumerate(persons, start=1):
        sheet.append([index, *person.display_row()])
        row_idx = index + 1
        for col in TEXT_COLUMNS:
            cell = sheet.cell(row=row_idx, column=col)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [6, 16, 14, 28, 14, 12, 36, 14, 22, 22, 24, 24, 18, 16, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        for cell in sheet[get_column_letter(index)]:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
    workbook.save(destination)
    workbook.close()
    return destination
